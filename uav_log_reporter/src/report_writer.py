"""Word(.docx) 보고서 작성 모듈."""
from __future__ import annotations

from pathlib import Path
from typing import Dict

from docx import Document
from docx.shared import Inches

from .metrics import FlightMetrics
from .utils import sec_to_kor_time



def _add_info_table(doc: Document, info: Dict[str, str]) -> None:
    table = doc.add_table(rows=0, cols=2)
    table.style = "Light Grid Accent 1"
    for k, v in info.items():
        row = table.add_row().cells
        row[0].text = k
        row[1].text = v



def write_report(
    output_docx: Path,
    day_no: int,
    date_str: str,
    flight_no: int,
    purpose: str,
    basic_info: Dict[str, str],
    plot_paths: Dict[str, Path],
    summaries: Dict[str, str],
    special_notes: str,
) -> None:
    """요구 포맷으로 docx 생성."""
    output_docx.parent.mkdir(parents=True, exist_ok=True)

    doc = Document()
    doc.add_heading(f"Day #{day_no} _ {date_str}", level=1)
    p = doc.add_paragraph()
    run = p.add_run(f"FLIGHT #{flight_no}")
    run.bold = True

    doc.add_paragraph("[비행 기본정보]")
    info = {
        "1. 비행회차": f"Flight #{flight_no}",
        "2. 비행 목적": purpose,
        "3. 총 비행시간": basic_info.get("total_time", "계산 불가"),
        "4. 총 비행거리": basic_info.get("total_distance", "계산 불가"),
        "5. 배터리 소모": basic_info.get("battery", "데이터 없음"),
        "6. 비행특이사항": special_notes,
    }
    _add_info_table(doc, info)

    section_info = [
        ("1) 고도", "altitude"),
        ("2) Roll", "roll"),
        ("3) Pitch", "pitch"),
        ("4) 배터리(Voltage)", "voltage"),
        ("5) 배터리(Current)", "current"),
    ]

    for title, key in section_info:
        doc.add_paragraph()
        h = doc.add_paragraph(title)
        h.runs[0].bold = True

        img = plot_paths.get(key)
        if img and img.exists():
            doc.add_picture(str(img), width=Inches(6.6))
        else:
            doc.add_paragraph("[그래프 데이터 없음]")

        doc.add_paragraph(f"<{summaries.get(key, '데이터 없음')}>")

    doc.save(output_docx)
